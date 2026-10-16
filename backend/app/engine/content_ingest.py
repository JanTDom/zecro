"""Smart Content Ingest & Editorial Redactor.

Enables importing raw content from web links (URL) or documents (.docx, .pdf, .txt),
extracting photographs, analyzing subject segmentation (AI cutout/masks), and
redacting professional Polish editorial newspaper copy (kicker, headline, deck, lead, body).
"""
from __future__ import annotations

import io
import mimetypes
import os
import re
import shutil
import urllib.parse
import uuid
from pathlib import Path
from typing import Any, Optional

import requests
from PIL import Image

from ..config import settings
from ..models.database import db
from ..models.schemas import Asset
from ..storage.asset_store import persist_asset
from ..storage.supabase_storage import get_storage_public_url, parse_storage_reference
from ..engine.polish_typography import polish_microtypography
from ..engine.image_ops import extract_alpha_contour
from ..engine.segmentation import segment_image
from ..ai.photo_focal import analyse_photo_focal
from ..ai.runtime import run_task
from ..ai.gemini_client import AIUnavailable

USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)


def download_and_extract_url(url: str) -> dict[str, Any]:
    """Fetch an internet article URL, extracting its title, lead, body text, and main photo URL."""
    url = url.strip()
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url

    headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
        "Accept-Language": "pl,en-US;q=0.7,en;q=0.3",
    }
    resp = requests.get(url, headers=headers, timeout=12)
    resp.raise_for_status()

    # Determine charset if specified
    encoding = resp.encoding if resp.encoding and resp.encoding.lower() != "iso-8859-1" else resp.apparent_encoding or "utf-8"
    html_content = resp.content.decode(encoding, errors="replace")

    domain = urllib.parse.urlparse(url).netloc.removeprefix("www.")

    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html_content, "html.parser")
    except ImportError:
        # Fallback regex parser if bs4 was missing
        return _regex_fallback_html_extract(html_content, domain, url)

    # 1. Clean out clutter
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "noscript", "svg", "form"]):
        tag.decompose()

    # 2. Extract Title
    title = ""
    og_title = soup.find("meta", property="og:title") or soup.find("meta", attrs={"name": "twitter:title"})
    if og_title and og_title.get("content"):
        title = og_title["content"].strip()
    if not title:
        h1 = soup.find("h1")
        if h1 and h1.get_text(strip=True):
            title = h1.get_text(strip=True)
    if not title and soup.title:
        title = soup.title.get_text(strip=True)

    # Clean domain from title if appended (e.g. "Tytuł artykułu - Onet Wiadomości")
    if " - " in title:
        parts = title.rsplit(" - ", 1)
        if len(parts[0]) > 15:
            title = parts[0]
    elif " | " in title:
        parts = title.rsplit(" | ", 1)
        if len(parts[0]) > 15:
            title = parts[0]

    # 3. Extract Description / Lead
    description = ""
    og_desc = soup.find("meta", property="og:description") or soup.find("meta", attrs={"name": "description"})
    if og_desc and og_desc.get("content"):
        description = og_desc["content"].strip()

    # 4. Extract Main Photo URL
    image_url = ""
    og_img = soup.find("meta", property="og:image") or soup.find("meta", attrs={"name": "twitter:image"})
    if og_img and og_img.get("content"):
        raw_img_url = og_img["content"].strip()
        image_url = urllib.parse.urljoin(url, raw_img_url)

    if not image_url:
        # Find first meaningful img in article
        for img in soup.find_all("img"):
            src = img.get("src") or img.get("data-src")
            if src and not any(skip in src.lower() for skip in ["logo", "icon", "avatar", "pixel", "banner"]):
                image_url = urllib.parse.urljoin(url, src)
                break

    # 5. Extract Article Body Text
    article_container = (
        soup.find("article")
        or soup.find(role="main")
        or soup.find(class_=re.compile(r"article|entry|post|content|story", re.I))
        or soup
    )
    paragraphs = []
    for p in article_container.find_all("p"):
        p_text = p.get_text(separator=" ", strip=True)
        # Avoid short noise lines, copyright or cookie warnings
        if len(p_text) > 35 and not any(noise in p_text.lower() for noise in ["ciasteczk", "cookie", "regulamin", "wszelkie prawa"]):
            paragraphs.append(p_text)

    body_text = "\n\n".join(paragraphs) if paragraphs else description

    return {
        "source_url": url,
        "domain": domain,
        "raw_title": title,
        "raw_description": description,
        "raw_body": body_text,
        "image_url": image_url,
    }


def _regex_fallback_html_extract(html: str, domain: str, base_url: str) -> dict[str, Any]:
    """Fallback text extractor if bs4 is not available."""
    title_m = re.search(r'<meta[^>]+property=["\']og:title["\'][^>]+content=["\']([^"\']+)["\']', html, re.I)
    title = title_m.group(1) if title_m else ""
    if not title:
        t_tag = re.search(r'<title[^>]*>(.*?)</title>', html, re.I | re.S)
        title = t_tag.group(1).strip() if t_tag else ""

    desc_m = re.search(r'<meta[^>]+property=["\']og:description["\'][^>]+content=["\']([^"\']+)["\']', html, re.I)
    desc = desc_m.group(1) if desc_m else ""

    img_m = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)["\']', html, re.I)
    image_url = img_m.group(1) if img_m else ""
    if image_url:
        image_url = urllib.parse.urljoin(base_url, image_url)

    clean_text = re.sub(r'<script.*?</script>', '', html, flags=re.I | re.S)
    clean_text = re.sub(r'<style.*?</style>', '', clean_text, flags=re.I | re.S)
    clean_text = re.sub(r'<[^>]+>', ' ', clean_text)
    clean_text = re.sub(r'\s+', ' ', clean_text).strip()

    return {
        "source_url": base_url,
        "domain": domain,
        "raw_title": title,
        "raw_description": desc,
        "raw_body": clean_text[:4000],
        "image_url": image_url,
    }


def extract_document_file(file_path: Path, filename: str) -> dict[str, Any]:
    """Extract editorial text and embedded images from uploaded DOCX, PDF, RTF, or TXT."""
    ext = file_path.suffix.lower()
    paragraphs: list[str] = []
    extracted_image_bytes: Optional[bytes] = None
    image_ext: str = ".jpg"

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    paragraphs.append(t)
            for table in doc.tables:
                for row in table.rows:
                    r_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if r_text:
                        paragraphs.append(r_text)

            # Check for embedded media images in document parts
            for rel in doc.part.related_parts.values():
                if "image" in getattr(rel, "content_type", ""):
                    extracted_image_bytes = rel.blob
                    image_ext = mimetypes.guess_extension(rel.content_type) or ".jpg"
                    break
        except Exception:
            # Fallback zip extract for docx
            import zipfile
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(file_path) as z:
                if "word/document.xml" in z.namelist():
                    tree = ET.fromstring(z.read("word/document.xml"))
                    for elem in tree.iter():
                        if elem.tag.endswith("}t") and elem.text:
                            paragraphs.append(elem.text)
                for item in z.namelist():
                    if item.startswith("word/media/") and any(item.lower().endswith(ie) for ie in [".jpg", ".jpeg", ".png", ".webp"]):
                        extracted_image_bytes = z.read(item)
                        image_ext = Path(item).suffix
                        break

    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        for page in reader.pages:
            t = page.extract_text()
            if t:
                paragraphs.append(t)
        # Attempt extracting images from pdf pages
        try:
            for page in reader.pages:
                if page.images:
                    extracted_image_bytes = page.images[0].data
                    image_ext = Path(page.images[0].name).suffix or ".jpg"
                    break
        except Exception:
            pass

    else:
        # Plain text / markdown / rtf
        raw = file_path.read_bytes()
        for enc in ["utf-8", "cp1250", "iso-8859-2", "latin1"]:
            try:
                text = raw.decode(enc)
                paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
                break
            except UnicodeDecodeError:
                continue

    full_text = "\n\n".join(paragraphs)
    raw_title = paragraphs[0] if paragraphs else Path(filename).stem

    return {
        "filename": filename,
        "raw_title": raw_title[:120],
        "raw_body": full_text,
        "extracted_image_bytes": extracted_image_bytes,
        "image_ext": image_ext,
    }


def process_editorial_photo(
    image_source: str | bytes,
    auto_cutout: bool = True,
    caption: str = "",
    credit: str = "",
) -> Optional[dict[str, Any]]:
    """Save an image asset, compute focal point, and optionally generate an AI cutout silhouette."""
    settings.UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    asset_id = str(uuid.uuid4())

    temp_path: Optional[Path] = None
    ext = ".jpg"

    if isinstance(image_source, str) and (image_source.startswith("http://") or image_source.startswith("https://")):
        try:
            resp = requests.get(image_source, headers={"User-Agent": USER_AGENT}, timeout=15)
            resp.raise_for_status()
            data = resp.content
            # Detect extension
            content_type = resp.headers.get("Content-Type", "")
            ext = mimetypes.guess_extension(content_type.split(";")[0]) or ".jpg"
        except Exception:
            return None
    elif isinstance(image_source, bytes):
        data = image_source
    else:
        return None

    if len(data) < 2000:
        return None  # Image too small to be meaningful editorial content

    file_name = f"editorial_{asset_id}{ext}"
    target_path = settings.UPLOADS_DIR / file_name
    target_path.write_bytes(data)

    width_px = height_px = None
    dpi = 300.0
    try:
        with Image.open(target_path) as img:
            width_px, height_px = img.size
            if img.info.get("dpi"):
                dpi = float(img.info["dpi"][0])
    except Exception:
        pass

    focal_review = analyse_photo_focal(target_path)
    contour = None
    alpha_mask_url = None
    has_cutout = False

    if auto_cutout:
        if ext.lower() == ".png":
            contour = extract_alpha_contour(str(target_path), 100.0, 100.0)
            if contour:
                has_cutout = True
        elif not (focal_review.get("status") == "ready" and focal_review.get("confidence", 0) >= 0.7
                  and focal_review.get("subject_kind") in {"scene"}):
            mask_out = settings.UPLOADS_DIR / f"mask_{asset_id}.png"
            seg_res = segment_image(target_path, mask_out, target_w_mm=100.0, target_h_mm=100.0)
            if seg_res.contour and seg_res.mask_path:
                contour = seg_res.contour
                has_cutout = True
                mask_ref = persist_asset(Path(seg_res.mask_path), prefix="uploads", key=Path(seg_res.mask_path).name)
                if mask_ref:
                    alpha_mask_url = (
                        get_storage_public_url(mask_ref)
                        if parse_storage_reference(mask_ref)
                        else f"/static/uploads/{Path(mask_ref).name}"
                    )

    stored_reference = persist_asset(target_path, prefix="uploads", key=target_path.name)
    if not stored_reference:
        stored_reference = f"/static/uploads/{target_path.name}"

    public_url = (
        get_storage_public_url(stored_reference)
        if parse_storage_reference(stored_reference)
        else f"/static/uploads/{target_path.name}"
    )

    focal = list(focal_review.get("focal_point") or [0.5, 0.5])
    aspect_ratio = (width_px / height_px) if width_px and height_px else 1.5

    # Register in DB
    asset = Asset(
        id=asset_id,
        filename=file_name,
        file_path=stored_reference,
        file_type="image",
        mime_type=mimetypes.guess_type(target_path.name)[0] or "image/jpeg",
        file_size_bytes=len(data),
        width=width_px,
        height=height_px,
        metadata={
            "focal_point": focal,
            "focal_status": focal_review.get("status", "unavailable"),
            "subject_kind": focal_review.get("subject_kind", "uncertain"),
            "aspect_ratio": aspect_ratio,
            "has_alpha_mask": has_cutout,
            "alpha_mask_url": alpha_mask_url,
            "contour_polygon": contour,
        }
    )
    try:
        db.save_asset(asset)
    except Exception:
        pass

    return {
        "id": f"photo-{asset_id[:8]}",
        "url": public_url,
        "previewUrl": alpha_mask_url or public_url,
        "caption": caption or "",
        "credit": credit or "Foto",
        "importance": 5,
        "autoCutout": auto_cutout,
        "hasCutoutMask": has_cutout,
        "alphaMaskUrl": alpha_mask_url,
        "cutoutContourPoints": contour,
        "aspectRatio": aspect_ratio,
        "placementHint": "lead-hero",
        "allowTextWrap": True,
        "widthPx": width_px,
        "heightPx": height_px,
        "dpi": dpi,
        "focalPoint": focal,
        "focalStatus": focal_review.get("status", "ready"),
        "focalSubjectKind": focal_review.get("subject_kind", "person" if has_cutout else "scene"),
    }


def redact_editorial_article(
    raw_title: str,
    raw_text: str,
    domain: str = "",
    project_id: Optional[str] = None,
) -> dict[str, Any]:
    """Redact raw source text into a ready-to-typeset Polish newspaper article.

    Uses AI (Gemini/OpenAI) behind strict journalistic and microtypographic standards,
    with a deterministic algorithmic fallback when AI is unavailable.
    """
    clean_raw = re.sub(r"\s+", " ", raw_text).strip()
    clean_title = re.sub(r"\s+", " ", raw_title).strip()

    # Try AI Redactor first
    prompt_system = (
        "Jesteś doświadczonym redaktorem naczelnym polskiej gazety (dziennik / magazyn). "
        "Twoim zadaniem jest zredagowanie surowego materiału w gotowy, profesjonalny artykuł prasowy.\n"
        "ZASADY OBOWIĄZUJĄCE:\n"
        "1. KICKER (nadtytuł): 1-3 słowa, wersaliki (np. RAPORT, GOSPODARKA, ŚWIAT, WYDARZENIA, NAUKA).\n"
        "2. HEADLINE (nagłówek): mocny, chwytliwy, uderzający polski tytuł prasowy. Prawdziwy język dziennikarski, "
        "bez sztuczności, bez pustosłowia.\n"
        "3. DECK (podtytuł): 1 zwięzłe zdanie wyjaśniające istotę sprawy.\n"
        "4. LEAD: 1-2 zdania wprowadzające w sedno wydarzeń, konkretne fakty i liczby.\n"
        "5. BODY: właściwy artykuł prasowy o długości 120-220 słów. Podzielony na 2-3 zwięzłe akapity rozdzielone pustą linią. "
        "Płynny, naturalny język polski, bez żargonu i bez klisz AI.\n"
        "6. PULL_QUOTE: jedno kluczowe, wyraziste zdanie lub cytat z tekstu.\n"
        "7. BYLINE: autor lub źródło (np. 'opr. Redakcja' lub domena źródła).\n"
        "Zwróć wyłącznie prawidłowy JSON ze strukturą: "
        '{"kicker": "...", "headline": "...", "deck": "...", "lead": "...", "body": "...", "pull_quote": "...", "byline": "..."}'
    )

    user_input = (
        f"Surowy tytuł: {clean_title}\n"
        f"Źródło / domena: {domain or 'Internet'}\n"
        f"Treść materiału źródłowego:\n{clean_raw[:4500]}"
    )

    try:
        task_res = run_task(
            "copy_fit",
            prompt_system,
            user_input,
            schema=None,
            project_id=project_id,
            temperature=0.4,
            max_tokens=1000,
        )
        data = task_res.data
        if isinstance(data, dict) and data.get("headline") and data.get("body"):
            return {
                "kicker": polish_microtypography((data.get("kicker") or "AKTUALNOŚCI").upper()),
                "headline": polish_microtypography(data.get("headline", clean_title).upper()),
                "deck": polish_microtypography(data.get("deck") or ""),
                "lead": polish_microtypography(data.get("lead") or ""),
                "body": polish_microtypography(data.get("body") or ""),
                "pull_quote": polish_microtypography(data.get("pull_quote") or ""),
                "byline": data.get("byline") or (f"opr. / {domain}" if domain else "Redakcja"),
                "provenance": task_res.provenance,
            }
    except Exception:
        # Fallback to deterministic editorial heuristics
        pass

    # Deterministic Algorithmic Redaction Fallback
    words = clean_raw.split()
    headline = clean_title if clean_title else " ".join(words[:8]).upper()
    if len(headline) > 90:
        headline = headline[:88].rsplit(" ", 1)[0] + "…"

    # Extract sentences
    sentences = re.split(r"(?<=[.!?])\s+", clean_raw)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 15]

    lead = sentences[0] if sentences else ""
    deck = sentences[1] if len(sentences) > 1 else ""
    body_sentences = sentences[2:10] if len(sentences) > 2 else sentences
    body = "\n\n".join([
        " ".join(body_sentences[:3]),
        " ".join(body_sentences[3:]),
    ]).strip() if len(body_sentences) >= 3 else " ".join(body_sentences)

    pull_quote = sentences[2] if len(sentences) > 2 else (sentences[0] if sentences else "")

    return {
        "kicker": "WYDANIE GŁÓWNE",
        "headline": polish_microtypography(headline.upper()),
        "deck": polish_microtypography(deck),
        "lead": polish_microtypography(lead),
        "body": polish_microtypography(body if body else clean_raw[:800]),
        "pull_quote": polish_microtypography(pull_quote[:120]),
        "byline": f"opr. / {domain}" if domain else "Redakcja",
        "provenance": {"provider": "deterministic_fallback", "model": "rule_based"},
    }
